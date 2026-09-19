import io
import re
from pathlib import Path
from urllib.parse import urlparse

import httpx
from bs4 import BeautifulSoup

TEXT_EXTS = {
    ".txt", ".md", ".csv", ".json", ".yaml", ".yml", ".xml", ".html", ".htm",
    ".js", ".jsx", ".ts", ".tsx", ".py", ".java", ".go", ".rb", ".php", ".cs",
}
SUPPORTED_EXTS = TEXT_EXTS | {".pdf", ".docx"}
MAX_FETCH_BYTES = 10 * 1024 * 1024


class LoaderError(Exception):
    """A problem with the document itself (shown to the user)."""


def _tidy(text: str, collapse_spaces: bool) -> str:
    text = text.replace("\x00", "").replace("\r\n", "\n").replace("\r", "\n")
    if collapse_spaces:
        text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _decode(data: bytes) -> str:
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return data.decode("latin-1", errors="replace")


def html_to_text(html: str) -> str:
    soup = BeautifulSoup(html, "html.parser")
    for tag in soup(["script", "style", "noscript", "svg", "nav", "footer"]):
        tag.decompose()
    return _tidy(soup.get_text("\n"), collapse_spaces=True)


def _load_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception as e:  # noqa: BLE001
            raise LoaderError("This PDF is password protected.") from e
    pages = [(p.extract_text() or "") for p in reader.pages]
    text = _tidy("\n\n".join(pages), collapse_spaces=True)
    if not text:
        raise LoaderError(
            "No text found in this PDF. It may be a scanned image (OCR is not supported yet)."
        )
    return text


def _load_docx(data: bytes) -> str:
    from docx import Document

    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    # Tables are appended after the paragraphs (original position is not kept)
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    text = _tidy("\n\n".join(parts), collapse_spaces=True)
    if not text:
        raise LoaderError("No text found in this DOCX file.")
    return text


def load_file(filename: str, data: bytes) -> str:
    ext = Path(filename).suffix.lower()
    if ext not in SUPPORTED_EXTS:
        raise LoaderError(f"Unsupported file type '{ext}'.")
    try:
        if ext == ".pdf":
            return _load_pdf(data)
        if ext == ".docx":
            return _load_docx(data)
        if ext in {".html", ".htm"}:
            text = html_to_text(_decode(data))
        else:
            # Keep indentation for code, YAML and JSON
            text = _tidy(_decode(data), collapse_spaces=False)
    except LoaderError:
        raise
    except Exception as e:  # noqa: BLE001
        raise LoaderError(f"Could not read this {ext} file: {e}") from e

    if not text:
        raise LoaderError("The file is empty.")
    return text


async def fetch_url(url: str) -> str:
    parsed = urlparse(url)
    if parsed.scheme not in {"http", "https"} or not parsed.netloc:
        raise LoaderError("Only http and https URLs are supported.")

    try:
        async with httpx.AsyncClient(
            follow_redirects=True,
            timeout=httpx.Timeout(25.0, connect=8.0),
            headers={"User-Agent": "Mozilla/5.0 (compatible; AITestingEngineBot/1.0)"},
        ) as client:
            r = await client.get(url)
    except httpx.HTTPError as e:
        raise LoaderError(f"Could not fetch the URL: {e}") from e

    if r.status_code >= 400:
        raise LoaderError(f"The URL returned HTTP {r.status_code}.")
    if len(r.content) > MAX_FETCH_BYTES:
        raise LoaderError("The page is larger than 10 MB.")

    ctype = r.headers.get("content-type", "").lower()
    try:
        if "application/pdf" in ctype:
            return _load_pdf(r.content)
        if "html" in ctype or not ctype:
            text = html_to_text(r.text)
            if len(text) < 80:
                raise LoaderError(
                    "This page has almost no readable text. It is probably rendered by "
                    "JavaScript. Upload the content as a file instead."
                )
            return text
        return _tidy(r.text, collapse_spaces=False)
    except LoaderError:
        raise
    except Exception as e:  # noqa: BLE001
        raise LoaderError(f"Could not read the page content: {e}") from e